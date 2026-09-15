#!/usr/bin/env python3
"""Generate the editable bilingual ROS Car/UWB PowerPoint deck."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "slides" / "ros_car_uwb_bilingual.md"
OUTPUT = ROOT / "slides" / "ros_car_uwb_bilingual.pptx"
NAVY = RGBColor(12, 28, 48)
CYAN = RGBColor(38, 198, 218)
WHITE = RGBColor(246, 249, 252)
MUTED = RGBColor(172, 190, 205)


def parse_slides(source: Path) -> tuple[str, list[tuple[str, list[str]]]]:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = next(line.removeprefix("title: ") for line in lines if line.startswith("title: "))
    slides: list[tuple[str, list[str]]] = []
    current_title: str | None = None
    bullets: list[str] = []
    for line in lines:
        if line.startswith("# "):
            if current_title:
                slides.append((current_title, bullets))
            current_title, bullets = line[2:], []
        elif current_title and line.startswith("- "):
            bullets.append(line[2:])
    if current_title:
        slides.append((current_title, bullets))
    return title, slides


def add_textbox(slide, x, y, w, h, text, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Noto Sans CJK SC"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def build(source: Path, output: Path) -> None:
    deck_title, slides = parse_slides(source)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = NAVY
    add_textbox(cover, 0.8, 1.25, 11.8, 2.1, deck_title, 30, WHITE, True)
    add_textbox(cover, 0.82, 3.65, 10.8, 0.6, "设计与验证计划 / Design & validation plan", 20, CYAN, True)
    add_textbox(cover, 0.82, 4.45, 10.8, 0.8, "非实现声明 · Not an implementation claim\n2026-09-15", 15, MUTED)

    for index, (title, bullets) in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = NAVY
        add_textbox(slide, 0.65, 0.38, 11.8, 0.72, title, 25, WHITE, True)
        accent = slide.shapes.add_shape(1, Inches(0.67), Inches(1.17), Inches(1.15), Inches(0.06))
        accent.fill.solid(); accent.fill.fore_color.rgb = CYAN; accent.line.fill.background()
        y = 1.45
        font_size = 17 if len(bullets) <= 5 else 15
        for bullet in bullets:
            add_textbox(slide, 0.86, y, 11.65, 0.62, f"•  {bullet}", font_size, WHITE)
            y += 0.83 if len(bullets) <= 5 else 0.72
        footer = add_textbox(slide, 0.72, 7.08, 11.9, 0.22, f"Lunar PNT · ROS Car / UWB     {index:02d}", 9, MUTED)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
